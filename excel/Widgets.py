from abc import ABC, abstractmethod

from openpyxl import chart
from openpyxl.chart import Reference
from openpyxl.utils import coordinate_to_tuple, get_column_letter


class Widget(ABC):
    def write(self, worksheet, cell='A1', offset_col=0, offset_row=0):
        row, col = coordinate_to_tuple(cell)
        cell = get_column_letter(col + offset_col) + str(row + offset_row)

        self._write(worksheet, cell)

    @abstractmethod
    def _write(self, worksheet, cell):
        pass

    @abstractmethod
    def update(self):
        pass

    @property
    @abstractmethod
    def size(self):
        pass


class Table(Widget):
    def __init__(self):
        self.header = []
        self.rows = []

    def _write(self, worksheet, cell):
        row, col = coordinate_to_tuple(cell)
        size_x, size_y = self.size
        rows = list(worksheet.get_squared_range(col, row, col + size_x - 1, row + size_y - 1))
        for i, header_cell in enumerate(rows[0]):
            header_cell.value = self.header[i]
        for i, row in enumerate(rows[1:]):
            for j, c in enumerate(row):
                c.value = self.rows[i][j]

    def append(self, *args):
        self.rows.append(args)

    @property
    def size(self):
        y = len(self.rows)
        x = 0 if y == 0 else len(self.rows[0])
        return x, y + 1


class PieChart(Table):
    def __init__(self, title):
        super().__init__()
        self.title = title
        self.header = ['key', 'value']

    def _write(self, worksheet, cell):
        super(PieChart, self)._write(worksheet, cell)

        row, col = coordinate_to_tuple(cell)

        pie = chart.PieChart()
        labels = Reference(worksheet, min_col=col, min_row=row + 1, max_row=row + len(self.rows))
        data = Reference(worksheet, min_col=col + 1, min_row=row, max_row=row + len(self.rows))
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(labels)
        pie.title = self.title
        worksheet.add_chart(pie, cell)
